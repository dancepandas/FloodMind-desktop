import json
from pathlib import Path


import json
from pathlib import Path

from floodmind_desktop_sidecar.agent import StubAgent
from floodmind_desktop_sidecar.app import create_app


def test_permission_respond_requires_ask_id(client):
    r = client.post("/api/permission/respond", json={"approved": True})
    assert r.status_code == 400


def test_permission_respond_ok(client):
    # StubAgent.respond_to_permission_ask always returns {"ok": True}.
    r = client.post("/api/permission/respond", json={"ask_id": "a1", "approved": True})
    assert r.status_code == 200
    assert r.get_json() == {"ok": True}


def test_permission_respond_propagates_agent_error():
    class ErrAgent(StubAgent):
        def respond_to_permission_ask(self, ask_id, approved):
            return {"error": "未找到待处理的权限请求"}

    c = create_app(ErrAgent()).test_client()
    r = c.post("/api/permission/respond", json={"ask_id": "a1", "approved": False})
    assert r.status_code == 400
    assert r.get_json() == {"error": "未找到待处理的权限请求"}


def test_feedback_requires_description(client):
    r = client.post("/api/feedback", json={"description": "   "})
    assert r.status_code == 400
    assert "description" in r.get_json()["error"]


def test_feedback_writes_file(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    body = {
        "description": "Something went wrong",
        "include_logs": True,
        "contact": "user@example.com",
        "logs": [{"level": "error", "message": "boom", "timestamp": "2026-07-15T00:00:00Z"}],
    }
    r = client.post("/api/feedback", json=body)
    assert r.status_code == 200
    data = r.get_json()
    assert data["ok"] is True
    assert data["file"].endswith(".json")
    assert "path" not in data            # absolute path must NOT be leaked in the response

    files = list((tmp_path / ".floodmind" / "feedback").glob("feedback-*.json"))
    assert len(files) == 1
    payload = json.loads(files[0].read_text(encoding="utf-8"))
    assert payload["description"] == "Something went wrong"
    assert payload["contact"] == "user@example.com"
    assert payload["include_logs"] is True
    assert payload["logs"] == body["logs"]
    assert payload["workspace"] == str(tmp_path)


def test_feedback_without_workspace_uses_app_data_fallback(client, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)          # _app_data_dir() == cwd
    monkeypatch.delenv("FLOODMIND_DESKTOP_WORKSPACE", raising=False)
    r = client.post("/api/feedback", json={"description": "no workspace", "include_logs": False})
    assert r.status_code == 200
    data = r.get_json()
    assert data["ok"] is True
    assert data["file"].endswith(".json")
    files = list((tmp_path / "feedback").glob("feedback-*.json"))
    assert len(files) == 1


def test_feedback_scrubs_secrets_from_logs(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    body = {
        "description": "crash",
        "include_logs": True,
        "logs": [
            {"level": "error", "message": "API key sk-abcdefghijklmnopqrstuvwxyz leaked", "timestamp": "t"},
            {"level": "info", "message": "Authorization: Bearer mF_9.B5f-4.1JqM token", "timestamp": "t"},
            {"level": "log", "message": "api_key: SUPERSECRET1234567 here", "timestamp": "t"},
        ],
    }
    r = client.post("/api/feedback", json=body)
    assert r.status_code == 200
    files = list((tmp_path / ".floodmind" / "feedback").glob("feedback-*.json"))
    payload = json.loads(files[0].read_text(encoding="utf-8"))
    blob = json.dumps(payload["logs"])
    assert "sk-abcdefghijklmnopqrstuvwxyz" not in blob
    assert "mF_9.B5f-4.1JqM" not in blob
    assert "SUPERSECRET1234567" not in blob
    assert blob.count("[REDACTED]") >= 3


def test_feedback_rejects_overlong_description(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    r = client.post("/api/feedback", json={"description": "x" * 8001})
    assert r.status_code == 400


def test_read_workspace_file_serves_content(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "hello.txt").write_text("world", encoding="utf-8")
    r = client.get("/api/workspace/file?path=hello.txt")
    assert r.status_code == 200
    assert r.get_json()["content"] == "world"


def test_read_workspace_file_forbidden(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    r = client.get("/api/workspace/file?path=../etc/passwd")
    assert r.status_code == 403


def test_read_workspace_file_directory(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "sub").mkdir()
    r = client.get("/api/workspace/file?path=sub")
    assert r.status_code == 400


def test_read_workspace_file_serves_docx_bytes(client, tmp_path, monkeypatch):
    """Office files are served as raw base64 + their mime (client renders)."""
    import base64

    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    import docx
    d = docx.Document()
    d.add_paragraph("Hello WordPreview")
    d.save(str(tmp_path / "note.docx"))
    r = client.get("/api/workspace/file?path=note.docx")
    assert r.status_code == 200
    j = r.get_json()
    assert j["previewable"] is True
    assert j["mime"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    raw = base64.b64decode(j["content"])
    assert raw[:2] == b"PK"  # OOXML is a ZIP archive


def test_read_workspace_file_serves_xlsx_bytes(client, tmp_path, monkeypatch):
    import base64

    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "Alice"
    ws["B1"] = 95
    wb.save(str(tmp_path / "data.xlsx"))
    r = client.get("/api/workspace/file?path=data.xlsx")
    j = r.get_json()
    assert j["previewable"] is True
    assert j["mime"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert base64.b64decode(j["content"])[:2] == b"PK"


def test_read_workspace_file_serves_pptx_bytes(client, tmp_path, monkeypatch):
    import base64

    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    from pptx import Presentation
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[1])
    prs.save(str(tmp_path / "deck.pptx"))
    r = client.get("/api/workspace/file?path=deck.pptx")
    j = r.get_json()
    assert j["previewable"] is True
    assert j["mime"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert base64.b64decode(j["content"])[:2] == b"PK"


def test_read_workspace_file_serves_legacy_doc_bytes(client, tmp_path, monkeypatch):
    """Legacy .doc is served raw too; the client renderer fails and offers
    'open in system' (the backend no longer distinguishes legacy formats)."""
    import base64

    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "old.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    r = client.get("/api/workspace/file?path=old.doc")
    j = r.get_json()
    assert j["previewable"] is True
    assert j["mime"] == "application/msword"
    assert base64.b64decode(j["content"])[:4] == b"\xd0\xcf\x11\xe0"


def test_attachment_thumb_returns_base64_for_workspace_image(client, tmp_path, monkeypatch):
    # repo convention: shared `client` fixture + monkeypatched FLOODMIND_DESKTOP_WORKSPACE
    # (the brief's `client_with_ws` fixture maps to this pattern — see existing
    # /api/workspace/file tests above).
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "pic.png").write_bytes(b"\x89PNG\r\n\x1a\n")  # tiny png-ish bytes
    r = client.get("/api/attachment/thumb?path=pic.png")
    assert r.status_code == 200
    data = r.get_json()
    assert data["mime"] == "image/png"
    assert data["content"]  # base64 non-empty


def test_attachment_thumb_rejects_non_image(client, tmp_path, monkeypatch):
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "note.txt").write_text("hi")
    r = client.get("/api/attachment/thumb?path=note.txt")
    assert r.status_code == 400


def test_attachment_thumb_rejects_oversized(client, tmp_path, monkeypatch):
    from floodmind_desktop_sidecar.attachments import MAX_IMAGE_BYTES
    monkeypatch.setenv("FLOODMIND_DESKTOP_WORKSPACE", str(tmp_path))
    (tmp_path / "big.png").write_bytes(b"\0" * (MAX_IMAGE_BYTES + 1))
    r = client.get("/api/attachment/thumb?path=big.png")
    assert r.status_code == 413
